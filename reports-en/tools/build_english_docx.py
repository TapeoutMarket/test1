#!/usr/bin/env python3
"""Build the two complete English legacy Word reports from their Markdown sources.

Run with the Python runtime returned by Codex load_workspace_dependencies.
Uses python-docx only; no network, environment files, wallet keys, or chain calls.
Rendered QA is performed separately with the documents skill render_docx.py.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from io import BytesIO
import re
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE

BASE = Path(__file__).resolve().parents[1]
REPORTS = BASE / 'legacy' / 'reports'
NAMES = ('tapeout-complete-test-report', 'tapeout-redemption-test-report')
STAMP = datetime(2026, 10, 20, 0, 0, 0, tzinfo=timezone.utc)


def xml(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def runs(paragraph, text, *, table=False):
    """Render Markdown emphasis, inline code, and real external hyperlinks."""
    parts = re.split(r'(\[[^\]]+\]\([^\)]+\)|\*\*.+?\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)', part)
        if match:
            label, url = match.groups()
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            props.append(xml('w:color', val='1F4E79'))
            props.append(xml('w:u', val='single'))
            props.append(xml('w:sz', val=19 if table else 21))
            run.append(props)
            node = OxmlElement('w:t')
            node.text = label
            run.append(node)
            link.append(run)
            paragraph._p.append(link)
        elif part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.2 if table else 9.8)
        else:
            paragraph.add_run(part)


def table_widths(rows, total=7.1):
    cols = len(rows[0])
    if cols == 2:
        return [1.70, total - 1.70]
    if cols == 3:
        if rows[0][0] == 'Reward ID':
            return [0.8, 1.55, total - 2.35]
        # Amount tables keep the identifying column generous.
        if 'Raw' in ' '.join(rows[0]) or 'integer' in ' '.join(rows[0]):
            return [3.25, 1.45, total - 4.70]
        return [2.05, 1.65, total - 3.70]
    if cols == 4:
        if rows[0][0] == 'Round':
            return [0.8, 3.3, 1.5, 1.5]
        if rows[0][0] == 'Number':
            return [0.75, 3.0, 1.35, 2.0]
        return [2.1, 1.30, 1.35, total - 4.75]
    if cols == 5:
        return [1.9, 1.25, 1.3, 1.3, total - 5.75]
    return [1.45] + [(total - 1.45) / (cols - 1)] * (cols - 1)


def add_table(doc, rows):
    widths = table_widths(rows)
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(xml('w:' + edge, val='single', sz=4, color='D9D9D9'))
    props.append(borders)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(xml('w:cantSplit'))
        if index == 0:
            row._tr.get_or_add_trPr().append(xml('w:tblHeader', val=1))
        for col, value in enumerate(values):
            cell = row.cells[col]
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom'):
                margins.append(xml('w:' + edge, w=80, type='dxa'))
            for edge in ('left', 'right'):
                margins.append(xml('w:' + edge, w=95, type='dxa'))
            cp.append(margins)
            if index == 0:
                cp.append(xml('w:shd', fill='EDEDED'))
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles['Table Body']
            paragraph.paragraph_format.keep_with_next = False
            if col and re.fullmatch(r'[0-9. /–−]+', value.strip('* ')):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(paragraph, value, table=True)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
    # Breathing room after the grid without a large forced gap.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(2)
    p.add_run().font.size = Pt(2)


def build(name):
    source = REPORTS / (name + '.md')
    text = source.read_text(encoding='utf-8')
    doc = Document()
    # The bundled default template has a blue Title border; remove every
    # inherited paragraph border so headings remain plain black typography.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    for style_name, size in [('Normal', 11), ('Title', 23), ('Heading 1', 16), ('Heading 2', 12.5)]:
        style = doc.styles[style_name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.12
        if style_name != 'Normal':
            style.font.bold = style_name != 'Title'
            style.paragraph_format.space_before = Pt(12 if style_name != 'Title' else 0)
            style.paragraph_format.keep_with_next = True
    table_style = doc.styles.add_style('Table Body', 1)
    table_style.base_style = doc.styles['Normal']
    table_style.font.size = Pt(9.5)
    table_style.paragraph_format.space_after = Pt(0)
    table_style.paragraph_format.space_before = Pt(0)
    table_style.paragraph_format.line_spacing = 1.10
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                values = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in values):
                    rows.append(values)
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith('#'):
            match = re.match(r'^(#{1,6})\s+(.+)$', line)
            if match:
                level, label = len(match.group(1)), match.group(2)
                label = label.replace("Ledger at each round's end", 'Ledger at the end of each round').replace('Can 45 / 55', 'Can 45 and 55')
                label = re.sub(r'[^\w\s]', ' ', label)
                label = re.sub(r'\s+', ' ', label).strip()
                p = doc.add_paragraph(style='Title' if level == 1 else 'Heading ' + str(min(level - 1, 2)))
                runs(p, label)
                i += 1
                continue
        p = doc.add_paragraph()
        if line.startswith('- '):
            p.style = doc.styles['List Bullet']
            line = line[2:]
        runs(p, line)
        i += 1
    doc.core_properties.title = doc.paragraphs[0].text
    doc.core_properties.subject = 'TapeOut legacy BSC mainnet tests and evidence'
    doc.core_properties.author = 'TapeOut'
    doc.core_properties.last_modified_by = 'TapeOut'
    doc.core_properties.created = STAMP
    doc.core_properties.modified = STAMP
    doc.core_properties.comments = ''
    # Stable ZIP ordering and timestamps make rebuilding byte reproducible.
    buf = BytesIO()
    doc.save(buf)
    output = REPORTS / (name + '.docx')
    with zipfile.ZipFile(buf) as zsrc, zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for member in sorted(zsrc.namelist()):
            info = zipfile.ZipInfo(member, (2026, 10, 20, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, zsrc.read(member))
    print(output)


if __name__ == '__main__':
    for report in NAMES:
        build(report)
